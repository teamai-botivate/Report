"""CSV/XLSX export from real query result rows (never screenshot/re-rendered data).

PNG export is deliberately client-side only (html-to-image on the actual rendered
report DOM node) — there is no server PNG endpoint. See ARCHITECTURE.md.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def rows_to_csv_bytes(rows: list[dict[str, Any]], columns: list[str] | None = None) -> bytes:
    if not rows and not columns:
        return b""
    fieldnames = columns or list(rows[0].keys())
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8-sig")


def rows_to_xlsx_bytes(
    rows: list[dict[str, Any]],
    columns: list[str] | None = None,
    sheet_title: str = "Data",
) -> bytes:
    fieldnames = columns or (list(rows[0].keys()) if rows else [])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Data"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, field in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fieldnames, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(field))

    for col_idx, field in enumerate(fieldnames, start=1):
        max_len = max([len(str(field))] + [len(str(r.get(field, ""))) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
